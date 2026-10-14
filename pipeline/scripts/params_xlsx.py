"""[params] 参数表：给不写 JSON 的人用的输入面。

pipeline 读的是 params.json。一个中文键名、绝对路径、六年行情倍数、六项分析师
假设的 JSON，办公人员既产生不了也修不了，只能再找一次 AI。同样这些事实放进一张
他本来就会编辑的工作簿，由这个模块在两者之间转换。

转换是逐格直译的：工作簿是事实来源，JSON 是产物，不做任何推断。空格表示「没提供」，
一路以 null 传下去，绝不变成 0 —— 模板把空读成 0 且不报任何错，所以「空」和「真的
是零」必须在整条链路上保持可区分。

财报文件名相对于本工作簿所在目录（或 --reports-dir）解析，所以整包可以搬到别的
机器上，不需要改任何绝对路径。报告期留空时按封面自动识别。

用法：
    python scripts/params_xlsx.py --init 参数表.xlsx --reports-dir 财报
    python scripts/params_xlsx.py --to-json 参数表.xlsx --out params.json
"""
import argparse
import datetime
import io
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

SHEET_BASIC = "基本信息"
SHEET_MARKET = "行情数据"
SHEET_MULT = "历史估值倍数"
SHEET_ANALYST = "分析师假设"
SHEET_REPORTS = "财报文件"

BASIC = [("公司名称", "name"), ("股票代码", "code"),
         ("报告期", "report_period"), ("估值日期", "valuation_date")]
MARKET = [("总股本(百万股)", "total_shares_mn"),
          ("流通股本(百万股)", "float_shares_mn"),
          ("收盘价(元)", "close_price"),
          ("港元兑人民币", "hkd_cny"), ("美元兑人民币", "usd_cny")]
MULT_COLS = ["PE", "PS", "PB", "PCF"]
ANALYST = ["营业收入增长率(2026E)", "营业收入增长率(2027E)", "营业成本率(2026E)",
           "管理费用率(2026E)", "研发费用率(2026E)", "新增长期借款率(2026E)"]
ALLOW_MISSING_HISTORY = "历史估值倍数缺失时是否继续(是/否)"
# 年度数据 要六列，缺一列就让派生假设归零，而工作簿不报错。缺年份时有两个选择：
# 什么都不写（安全，但办公人员手上什么都没有），或把已抽到的事实写进去、文件名带上
# 待补数据、卡片写清哪些数字别用。默认选后者，因为「核对事实」本身有用，
# 而这份文件从来没有被当成可用的估值交付过。
ALLOW_INCOMPLETE = "历史年份不全时仍写出表格(是/否)"


def _years(end_year, n=6):
    return ["%d-12-31" % y for y in range(end_year - n + 1, end_year + 1)]


def _as_date(v):
    """A cell the user typed as a date, a string, or a serial -> YYYY-MM-DD."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v).strip().replace("/", "-").replace(".", "-")
    parts = s.split("-")
    if len(parts) == 3:
        try:
            return "%04d-%02d-%02d" % (int(parts[0]), int(parts[1]), int(parts[2]))
        except ValueError:
            return s
    return s


def _num(v):
    """Blank stays blank. A blank read as 0 is the failure this guards against."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("%", "")
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _yes(v):
    return str(v).strip().lower() in ("是", "y", "yes", "true", "1", "继续")


def _label_rows(ws, start=1):
    """{label: value} from a two-column 项目/值 sheet."""
    out = {}
    for r in range(start, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k is None or str(k).strip() == "":
            continue
        out[str(k).strip()] = ws.cell(r, 2).value
    return out


def basic_info(path):
    """(名称, 代码, 报告期) from 基本信息, blank fields as None.

    Only used to name the output folder, so it must not raise: a half-filled sheet
    has to reach the run that explains what is missing, and the pipeline validates
    the same sheet a moment later anyway.
    """
    try:
        rows = _label_rows(load_workbook(path, data_only=True)[SHEET_BASIC])
    except Exception:
        return None, None, None
    name = str(rows.get("公司名称") or "").strip()
    code = str(rows.get("股票代码") or "").strip()
    return name or None, code or None, _as_date(rows.get("报告期"))


def read_params(path, reports_dir=None):
    wb = load_workbook(path, data_only=True)
    sheet_dir = os.path.dirname(os.path.abspath(path))
    base = reports_dir or sheet_dir

    def sheet(name):
        if name not in wb.sheetnames:
            raise SystemExit("参数表缺少工作表「%s」" % name)
        return wb[name]

    b = _label_rows(sheet(SHEET_BASIC))
    m = _label_rows(sheet(SHEET_MARKET))

    company = {key: (str(b.get(cn)).strip() if b.get(cn) not in (None, "") else None)
               for cn, key in BASIC}
    company["report_period"] = _as_date(b.get("报告期"))
    company["valuation_date"] = _as_date(b.get("估值日期")) or company["report_period"]
    for req in ("name", "code", "report_period"):
        if not company.get(req):
            raise SystemExit("参数表「%s」没填：%s" % (SHEET_BASIC, req))

    market = {key: _num(m.get(cn)) for cn, key in MARKET}
    market["allow_missing_history"] = _yes(m.get(ALLOW_MISSING_HISTORY))
    allow_incomplete = _yes(b.get(ALLOW_INCOMPLETE))

    ws = sheet(SHEET_MULT)
    history = {}
    for r in range(2, ws.max_row + 1):
        d = _as_date(ws.cell(r, 1).value)
        if not d:
            continue
        row = {c: _num(ws.cell(r, 2 + i).value) for i, c in enumerate(MULT_COLS)}
        if any(v is not None for v in row.values()):
            history[d] = row
    market["history"] = history

    ws = sheet(SHEET_ANALYST)
    analyst = {}
    for r in range(2, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k is None or str(k).strip() == "":
            continue
        v = _num(ws.cell(r, 2).value)
        if v is not None:
            analyst[str(k).strip()] = v

    ws = sheet(SHEET_REPORTS)
    reports, missing = [], []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if name is None or str(name).strip() == "":
            continue
        name = str(name).strip()
        # The sheet records paths relative to itself, so it stays valid when the
        # folder moves. --reports-dir is a fallback for a sheet that lists bare
        # file names, never an override of a path that already resolves.
        if os.path.isabs(name):
            cands = [name]
        else:
            cands = [os.path.join(sheet_dir, name), os.path.join(base, name)]
        full = next((c for c in cands if os.path.exists(c)), None)
        if full is None:
            missing.append(name)
            continue
        reports.append({"pdf": full, "period": _as_date(ws.cell(r, 2).value)})
    if missing:
        raise SystemExit("参数表「%s」里这些文件找不到（相对目录 %s）：\n  %s"
                         % (SHEET_REPORTS, base, "\n  ".join(missing)))
    if not reports:
        raise SystemExit("参数表「%s」里一份 PDF 都没有。" % SHEET_REPORTS)

    # A report whose period was left blank has the document decide, exactly as the
    # preflight would: the cover states what it is, and a quarter typed into a
    # year-end column is the silent failure this whole pipeline exists to stop.
    import detect as _detect
    for rep in reports:
        if rep["period"]:
            continue
        rec = _detect.detect(rep["pdf"])
        if not rec.get("period"):
            raise SystemExit("认不出这份 PDF 的报告期，请在参数表里手填：%s"
                             % os.path.basename(rep["pdf"]))
        rep["period"] = rec["period"]
        rep["kind"] = rec["kind"]

    return {"company": company, "market": market,
            "reports": reports, "analyst": analyst,
            "allow_incomplete_history": allow_incomplete,
            "_comment": "由 scripts/params_xlsx.py 从参数表生成，不要手改；改参数表再转换。"}


def write_template(path, reports_dir=None, report_period="2025-12-31"):
    wb = Workbook()
    head = Font(bold=True)

    ws = wb.active
    ws.title = SHEET_BASIC
    ws.append(["项目", "值"])
    for cn, _ in BASIC:
        ws.append([cn, ""])
    ws.append([ALLOW_INCOMPLETE, "是"])
    ws["B4"] = report_period
    ws["B5"] = report_period
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 34
    for c in ("A1", "B1"):
        ws[c].font = head

    ws = wb.create_sheet(SHEET_MARKET)
    ws.append(["项目", "值"])
    for cn, _ in MARKET:
        ws.append([cn, ""])
    ws.append([ALLOW_MISSING_HISTORY, "否"])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    for c in ("A1", "B1"):
        ws[c].font = head

    ws = wb.create_sheet(SHEET_MULT)
    ws.append(["年度"] + MULT_COLS)
    for y in _years(int(str(report_period)[:4])):
        ws.append([y] + [""] * len(MULT_COLS))
    ws.column_dimensions["A"].width = 14
    for c in ("A1", "B1", "C1", "D1", "E1"):
        ws[c].font = head

    ws = wb.create_sheet(SHEET_ANALYST)
    ws.append(["项目", "值"])
    for k in ANALYST:
        ws.append([k, ""])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    for c in ("A1", "B1"):
        ws[c].font = head

    ws = wb.create_sheet(SHEET_REPORTS)
    ws.append(["PDF 文件名", "报告期（留空则按封面自动识别）"])
    found = []
    if reports_dir and os.path.isdir(reports_dir):
        found = sorted(f for f in os.listdir(reports_dir)
                       if f.lower().endswith(".pdf"))
    import detect as _detect
    # The path is written relative to the sheet, so the whole folder can be moved
    # to another machine (or another drive letter) without editing anything, and
    # so nobody has to remember to pass --reports-dir later.
    base = os.path.dirname(os.path.abspath(path))
    for f in found:
        full = os.path.join(reports_dir, f)
        try:
            shown = os.path.relpath(full, base)
        except ValueError:          # different drive: no relative form exists
            shown = full
        rec = _detect.detect(full)
        ws.append([shown, rec.get("period") or ""])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 32
    for c in ("A1", "B1"):
        ws[c].font = head

    wb.save(path)
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--init", metavar="XLSX", help="生成一张空白参数表")
    ap.add_argument("--to-json", metavar="XLSX", help="把参数表转成 params.json")
    ap.add_argument("--reports-dir", help="财报 PDF 所在目录（默认参数表所在目录）")
    ap.add_argument("--report-period", default="2025-12-31")
    ap.add_argument("--out", help="--to-json 的输出路径")
    a = ap.parse_args()

    if a.init:
        found = write_template(a.init, a.reports_dir, a.report_period)
        print("已生成参数表：%s" % a.init)
        if found:
            print("  已自动列出 %d 份财报，报告期按封面识别；识别不出的请手填" % len(found))
        print("  填好后运行：python scripts/params_xlsx.py --to-json \"%s\" --out params.json"
              % a.init)
        return 0

    if a.to_json:
        out = a.out or os.path.splitext(a.to_json)[0] + ".json"
        p = read_params(a.to_json, a.reports_dir)
        json.dump(p, io.open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("公司：%s（%s）报告期 %s"
              % (p["company"]["name"], p["company"]["code"], p["company"]["report_period"]))
        print("财报 %d 份；历史倍数 %d 年；分析师假设 %d 项"
              % (len(p["reports"]), len(p["market"]["history"]), len(p["analyst"])))
        print("历史年份不全时仍写出表格：%s" % ("是" if p["allow_incomplete_history"] else "否"))
        blank = [k for k in ANALYST if k not in p["analyst"]]
        if blank:
            print("  还空着的分析师假设：%s" % "、".join(blank))
        print("wrote", out)
        return 0

    ap.error("需要 --init 或 --to-json")


if __name__ == "__main__":
    sys.exit(main())
